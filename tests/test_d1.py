"""D1: the final case folder (Markdown, Word, PDF, Excel, images; Turkish) as briefs, and no-OpenAPI discovery."""
import json
import shutil
from pathlib import Path
from urllib.parse import urlparse

import pytest

from builder import build as BD
from builder import doctext, sdk
from builder import main as M

from test_main_dry import HOST, ROOT, env, trace_rows, clean_env  # noqa: F401 - fixtures
from test_mock_url import MOCK, MOCK_HOST, urls

TR_PDF = {}
TR = "çğıİöşü ÇĞÖŞÜ Müşteri kartı iptali"


def make_docx(path):
    import docx
    d = docx.Document()
    d.add_heading("Hizmet Sözleşmesi", level=1)
    d.add_paragraph("Paragraf: " + TR)
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "Alan", "Açıklama"
    t.cell(1, 0).text, t.cell(1, 1).text = "müşteri_no", "Şube | kodu"
    d.add_paragraph("Tablodan sonra")
    d.save(str(path))


def make_xlsx(path):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Müşteriler"
    ws.append(["Ad", "Şehir"])
    ws.append(["Ayşe Öztürk", "İzmir"])
    ws2 = wb.create_sheet("Kodlar")
    ws2.append(["kod", "anlam"])
    ws2.append([42, "ığdır"])
    wb.save(str(path))


def make_pdf(path, pages):
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.pdfgen import canvas
    font = "Helvetica"
    for cand in ("/usr/share/fonts/TTF/DejaVuSans.ttf", "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
                 "/usr/share/fonts/noto/NotoSans-Regular.ttf", "/usr/share/fonts/TTF/NotoSans-Regular.ttf"):
        if Path(cand).exists():
            pdfmetrics.registerFont(TTFont("TRFont", cand))
            font = "TRFont"
            break
    c = canvas.Canvas(str(path), pagesize=A4)
    for text in pages:
        c.setFont(font, 12)
        c.drawString(72, 750, text)
        c.showPage()
    c.save()
    return font != "Helvetica"


@pytest.fixture
def case(tmp_path):
    d = tmp_path / "Banka Vakası"
    (d / "ekler").mkdir(parents=True)
    (d / "BRIEF.md").write_text("# Brief\n" + TR + "\n", encoding="utf-8")
    (d / "README.md").write_text("Belgeler ne?\n", encoding="utf-8")
    (d / "a-analiz.md").write_text("analiz\n", encoding="utf-8")
    (d / "duyurular.txt").write_bytes(("Duyuru: " + TR + "\n").encode("cp1254"))   # not utf-8
    (d / "test-musteriler.csv").write_text("ad,şehir\nAyşe,İzmir\n", encoding="utf-8")
    make_docx(d / "sözleşme.docx")
    make_xlsx(d / "ekler" / "müşteriler.xlsx")
    TR_PDF["ok"] = make_pdf(d / "soru-cevap.pdf", ["Birinci sayfa " + TR, "Ikinci sayfa"])
    (d / "akis.png").write_bytes(b"\x89PNG\r\n\x1a\n" + b"\x00" * 16)
    (d / "garip.bin").write_bytes(b"\x00\xff\x00\xfe")
    (d / ".gizli.md").write_text("hidden\n")
    (d / ".env").write_text("SECRET=1\n")
    return d


def test_docx_text_and_table(case):
    t = doctext.to_text(case / "sözleşme.docx")
    assert "## Hizmet Sözleşmesi" in t and TR in t
    assert "| Alan | Açıklama |" in t and "| müşteri_no | Şube \\| kodu |" in t
    assert t.index("Paragraf") < t.index("| Alan") < t.index("Tablodan sonra")


def test_xlsx_every_sheet(case):
    t = doctext.to_text(case / "ekler" / "müşteriler.xlsx")
    assert "## Sheet: Müşteriler" in t and "| Ayşe Öztürk | İzmir |" in t
    assert "## Sheet: Kodlar" in t and "| 42 | ığdır |" in t
    assert t.index("Müşteriler") < t.index("Kodlar")


def test_pdf_pages_in_order(case):
    t = doctext.to_text(case / "soru-cevap.pdf")
    assert "[page 1]" in t and "[page 2]" in t and t.index("Birinci") < t.index("Ikinci")
    if TR_PDF["ok"]:
        assert TR in t


def test_image_and_unknown_listed_never_crash(case, tmp_path):
    assert doctext.to_text(case / "akis.png") == "akis.png (image, not read)"
    assert "not read" in doctext.to_text(case / "garip.bin")
    bad = tmp_path / "bozuk.docx"
    bad.write_bytes(b"not a zip")
    assert "could not be read" in doctext.to_text(bad)


def test_encoding_fallback(case):
    assert TR in doctext.to_text(case / "duyurular.txt")


def test_order_brief_readme_then_rest_recursive(case):
    rels = [p.relative_to(case).as_posix() for p in doctext.case_files(case)]
    assert rels[:2] == ["BRIEF.md", "README.md"]
    assert "ekler/müşteriler.xlsx" in rels and ".gizli.md" not in rels and ".env" not in rels
    assert rels[2:] == sorted(rels[2:], key=str.lower)
    assert rels == [p.relative_to(case).as_posix() for p in doctext.case_files(case)]   # stable


def test_load_folder_headings(case):
    t = doctext.load_folder(case)
    assert t.startswith("# BRIEF.md\n") and "# ekler/müşteriler.xlsx" in t and "# akis.png\n\nakis.png (image, not read)" in t
    assert t.index("# BRIEF.md") < t.index("# README.md") < t.index("# a-analiz.md")


def test_workdir_gets_text_renderings(case, tmp_path, monkeypatch):
    monkeypatch.setattr(sdk, "WORK_ROOT", tmp_path / "work")
    shutil.copy(ROOT / "tools/openapi.json", case / "openapi.json")
    oa = json.loads((ROOT / "tools/openapi.json").read_text())
    work = sdk.make_workdir([str(case)], oa, "t")
    names = sdk.brief_names(work)
    assert names[:2] == ["BRIEF.md", "README.md"]
    assert "sözleşme.docx.md" in names and "ekler__müşteriler.xlsx.md" in names and "akis.png.md" in names
    assert "openapi.json" not in names and ".env" not in names
    assert (work / "sözleşme.docx.md").read_text(encoding="utf-8").startswith("# sözleşme.docx\n")
    text = sdk.brief_text(work)
    assert "| Ayşe Öztürk | İzmir |" in text and "| Alan | Açıklama |" in text and TR in text
    assert (work / "duyurular.txt").read_text(encoding="utf-8").count(TR) == 1   # cp1254 in, utf-8 out


# ------------------------------------------------------------------------------------------------ OpenAPI


def test_yaml_openapi_found_recursively_and_not_a_brief(case):
    oa = json.loads((ROOT / "tools/openapi.json").read_text())
    import yaml
    (case / "ekler" / "servis.yaml").write_text(yaml.safe_dump(oa, allow_unicode=True), encoding="utf-8")
    assert M.resolve_openapi([str(case)], None) == (case / "ekler" / "servis.yaml").resolve()
    assert sdk.load_openapi(case / "ekler" / "servis.yaml") == oa
    assert all(p.name != "servis.yaml" for p in sdk.expand_briefs([str(case)]))
    argv = BD.main_argv([str(case), MOCK])
    assert argv[argv.index("--openapi") + 1] == str((case / "ekler" / "servis.yaml").resolve())


def test_no_openapi_folder_does_not_crash_discovery(case):
    with pytest.raises(M.OpenapiNotFound) as e:
        M.resolve_openapi([str(case)], None)
    assert e.value.none
    argv = BD.main_argv([str(case), MOCK])
    assert argv == [str(case), "--openapi-fallback", "--slot", "final-banka-vakas", "--mock-url", MOCK]


def test_fallback_needs_mock_url(case):
    with pytest.raises(SystemExit):
        M.parse_args([str(case), "--openapi-fallback"])


def test_main_without_fallback_still_exits_2(env):
    (env.brief.parent / "openapi.json").unlink()
    assert M.main([str(env.brief), "--no-selftest", "--mock-url", MOCK, "--out-dir", str(env.out),
                   "--trace-dir", str(env.traces)]) == 2


def test_main_fallback_exports_mock_only_and_traces_loudly(env, capsys):
    (env.brief.parent / "openapi.json").unlink()
    rc = M.main([str(env.brief.parent), "--no-selftest", "--openapi-fallback", "--mock-url", MOCK,
                 "--out-dir", str(env.out), "--trace-dir", str(env.traces)])
    assert rc == 0
    rows = trace_rows(env)
    fb = [r for r in rows if r["kind"] == "openapi_fallback"]
    assert fb and fb[0]["openapi"] == str(M.KIT_OPENAPI) and fb[0]["mock_url"] == MOCK
    assert "NOT MACHINE-READABLE" in capsys.readouterr().err
    [(slot, text)] = env.saved
    assert HOST not in text and MOCK_HOST in text
    [exp] = list(env.out.glob("*.workflow.json"))
    assert {urlparse(u).netloc for u in urls(json.loads(exp.read_text())["workflow"])} == {MOCK_HOST}


# ------------------------------------------------------------------------------------------------ walker safety


def _rels(case):
    return [p.relative_to(case).as_posix() for p in doctext.case_files(case)]


def test_symlinks_not_followed(case, tmp_path):
    outside = tmp_path / "outside"
    outside.mkdir()
    (outside / "sır.md").write_text("DIŞARIDA\n")
    (case / "link.md").symlink_to(outside / "sır.md")
    (case / "linkdir").symlink_to(outside, target_is_directory=True)
    rels = _rels(case)
    assert "link.md" not in rels and not any(r.startswith("linkdir") for r in rels)
    assert "DIŞARIDA" not in doctext.load_folder(case)
    assert "not read" in doctext.to_text(case / "link.md")


def test_walk_stays_inside_root(case, tmp_path):
    for p in doctext.case_files(case):
        p.resolve().relative_to(case.resolve())   # raises if outside
    sub = case / "ekler"
    assert all(p.resolve().is_relative_to(sub.resolve()) for p in doctext.case_files(sub))


def test_dotfiles_and_secret_names(case):
    (case / ".git").mkdir()
    (case / ".git" / "config").write_text("x")
    for n in ("server.key", "cert.pem", "id_rsa", "credentials.json", "env.sh"):
        (case / n).write_text("TOPSECRET\n")
    rels = _rels(case)
    assert not any(r.startswith(".") for r in rels)
    t = doctext.load_folder(case)
    assert "TOPSECRET" not in t and "# server.key\n\nserver.key (secret-looking file, not read)" in t
    assert "id_rsa (secret-looking file, not read)" in t
